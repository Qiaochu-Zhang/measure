#!/usr/bin/env python3
"""Offline analytical, background, compatibility and export regression tests.

Generated images/exports live in a temporary directory. Use --output to retain
the numerical report. These tests do not establish real-SEM accuracy.
"""

from argparse import ArgumentParser, Namespace
from contextlib import redirect_stdout, redirect_stderr
from copy import deepcopy
from io import StringIO
import importlib.metadata
import json
from pathlib import Path
import platform
import tempfile
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

import sem_cd_measure_200k_batch_V1_15 as app
from cdsem_psd import segment_spectra, scalar_metrics
from cdsem_refinement import viterbi_path, fit_erf
from scipy.special import erf

BASE = Path(__file__).resolve().parent


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def analytical_tests():
    report = {}
    p = app.edges.MeasurementParams(pixel_size_nm=1, sample_number=128)
    rows = [
        dict(
            space_index=1,
            space_role="center",
            sample_index=i,
            sample_y=float(i),
            left_edge_x=20 + 0.1 * i,
            right_edge_x=80 + 0.1 * i,
            valid=True,
        )
        for i in range(128)
    ]
    metrics = app.coordinate_metrics_by_space(rows, p, 4)[1]
    require(
        abs(metrics["rotated_raw_CD_nm"] - 60 / np.sqrt(1.01)) < 1e-9,
        "PCA projected CD",
    )
    require(metrics["rotated_raw_LER_left_nm"] < 1e-10, "straight tilted edge LER")
    require(metrics["unrotated_raw_LER_left_nm"] > 10, "unrotated tilt retained")
    single = app.coordinate_metrics_by_space(deepcopy(rows), p, 128)[1]
    require(
        np.isnan(single["rotated_group4_LWR_nm"]), "single group must not report zero"
    )
    require(single["valid_group_count"] == 1, "group support count")
    one = app.coordinate_metrics_by_space([deepcopy(rows[0])], p, 1)[1]
    require(np.isnan(one["rotated_raw_LER_left_nm"]), "single edge cannot support LER")
    missing = deepcopy(rows)
    missing[2]["left_edge_x"] = np.nan
    missing[9]["right_edge_x"] = np.nan
    common = app.coordinate_metrics_by_space(missing, p, 4)[1]
    require(common["valid_sample_count"] == 126, "paired validity mask")
    report["geometry_and_support"] = "PASS"

    settings = Namespace(
        psd_gap_mode="segments",
        psd_max_gap=2,
        psd_min_segment=16,
        psd_window="boxcar",
        psd_detrend="constant",
    )
    wave = 2 * np.sin(2 * np.pi * np.arange(128) / 16)
    f, power, meta = segment_spectra(wave, 1.0, settings)[0]
    scalar = scalar_metrics(f, power, 100.0)
    require(abs(scalar["variance_psd_nm2"] - 2) < 1e-12, "PSD Parseval normalization")
    require(abs(f[np.argmax(power)] - 1 / 16) < 1e-12, "PSD physical frequency")
    alternating = np.tile([59.0, 61.0], 64)
    f, power, _ = segment_spectra(alternating, 1.0, settings)[0]
    require(
        abs(scalar_metrics(f, power, 100.0)["roughness_psd_3sigma_nm"] - 3) < 1e-10,
        "PSD must preserve alternating width; group4 would erase it",
    )
    gap = wave.copy()
    gap[50:60] = np.nan
    parts = segment_spectra(gap, 1.0, settings)
    require(
        len(parts) == 2 and [s[2]["used_slots"] for s in parts] == [50, 68],
        "separate missing-data segments",
    )
    gap[50:60] = wave[50:60]
    gap[50:52] = np.nan
    settings.psd_gap_mode = "interpolate"
    hard = np.zeros(128, bool)
    hard[50:52] = True
    require(
        len(segment_spectra(gap, 1.0, settings, hard_gaps=hard)) == 2,
        "background gaps must not be interpolated",
    )
    report["raw_psd_no_averaging"] = dict(
        status="PASS", known_variance=2.0, measured_variance=scalar["variance_psd_nm2"]
    )

    # Selected gradient peak is just inside the actual descending crossing.
    profile = np.r_[np.full(15, 220.0), np.linspace(220, 20, 7), np.full(25, 20.0)]
    legacy = app.edges.interpolate_threshold_crossing(profile, 150.0, 19, 35, "left")
    fixed = app.edges.bounded_threshold_crossing(
        profile, 150.0, 19, 35, "left", 12, 26, 18
    )
    require(legacy is None and fixed is not None, "outside-peak threshold recovery")
    require(
        app.edges.bounded_threshold_crossing(profile, 150.0, 19, 35, "left", 22, 26, 24)
        is None,
        "crossing must remain inside anchor window",
    )
    report["bounded_crossing"] = dict(status="PASS", fixed_position_px=fixed)

    layers = [
        [dict(left=0.0, right=60.0, cost=0.0)],
        [],
        [dict(left=1.0, right=61.0, cost=0.0)],
    ]
    path, score = viterbi_path(layers, 5.0, 3.0, 6.0)
    require(set(path) == {0, 2}, "Viterbi must retain a gap")
    x = np.arange(80.0)
    intensity = 100 - 80 * erf((x - 35.3) / (np.sqrt(2) * 1.5))
    fit, reason = fit_erf(intensity, 35.0, "left", 50.0, 8, 2.0, 0.2)
    require(
        fit is not None and abs(fit["x"] - 35.3) < 1e-3, "ERF subpixel fit: " + reason
    )
    report["viterbi_and_erf"] = "PASS"

    rng = np.random.default_rng(314)
    blanks = {
        "flat": np.full((256, 512), 127, np.uint8),
        "low_noise": np.clip(127 + rng.normal(0, 1, (256, 512)), 0, 255).astype(
            np.uint8
        ),
        "horizontal_bands": np.broadcast_to(
            ((np.arange(256) // 25) % 2 * 180 + 30)[:, None], (256, 512)
        ),
        "checkerboard": ((np.indices((256, 512)).sum(axis=0) // 3) % 2 * 180 + 30),
    }
    for name, gray in blanks.items():
        try:
            app.regions.find_pattern_region(gray, 60)
        except app.regions.NoPatternRegion:
            pass
        else:
            raise AssertionError("Background accepted: " + name)
    report["background_rejection"] = list(blanks)
    from check_localization import analytical_tests as locator_tests

    report["v115_localization"] = locator_tests(require, app)
    from check_line_pitch import analytical_tests as pitch_tests

    report["line_pitch"] = pitch_tests(require, app)
    return report


def integration_tests(temp):
    cases, cached = {}, {}

    def execute(name, root=None, extra=(), pattern="trench", reference=60):
        output = temp / name
        fixture = root or BASE / f"examples/input_{pattern}"
        args = [
            "--root",
            str(fixture),
            "--output",
            str(output),
            "--pattern",
            pattern,
            "--pixel-size",
            "1",
            *(
                [f"--{pattern}-reference-nm", str(reference)]
                if reference is not None
                else []
            ),
            "--max-number",
            "4",
            "--no-auto-machine-comparison",
            "--skip-statistics-plots",
            *extra,
        ]
        print("integration:", name, flush=True)
        log = StringIO()
        with redirect_stdout(log), redirect_stderr(log):
            rc = app.main(args)
        require(rc == 0, f"{name} returned {rc}\n" + log.getvalue())
        frame = pd.read_csv(output / "image_summary.csv")
        book = load_workbook(
            output / "CD_measurement_200K_V1_15_results.xlsx", read_only=True
        )
        require(
            book.sheetnames[0] == "measurement_summary",
            "summary must be first worksheet",
        )
        summary = book["measurement_summary"]
        require(
            next(summary.values)[-1] == "旋转_pitch_CD_nm",
            "pitch must be last even on failure/blank images",
        )
        require(
            [c.value for c in next(summary.iter_rows())][:3]
            == ["image", "status", "method"],
            "metadata first",
        )
        require(
            [summary.cell(i, 3).value for i in (2, 3, 4)] == ["mixed", "V13", "V10"],
            "mixed first",
        )
        book.close()
        case = dict(
            exit_code=rc,
            statuses=list(frame.status),
            image_rows=len(frame),
            settings=json.loads((output / "settings.json").read_text()),
        )
        cols = [
            "旋转_mixed_CD_nm",
            "旋转_mixed_LER_left_nm",
            "旋转_mixed_LER_right_nm",
            "旋转_mixed_LWR_nm",
        ]
        case["rotated_mixed"] = {
            c: (float(frame[c].iloc[0]) if pd.notna(frame[c].iloc[0]) else None)
            for c in cols
        }
        cases[name] = case
        return output, frame

    real_measure = app.edges.measure_image

    def capture(path, params, **kwargs):
        result = real_measure(path, params, **kwargs)
        if params.average_range_px == 32:
            cached[params.engine_kind] = deepcopy(result)
        return result

    with patch.object(app.edges, "measure_image", side_effect=capture):
        output, frame = execute("default_trench")
    require(frame.status.iloc[0] == "OK", "default fixture should no longer be REVIEW")
    samples = pd.read_csv(output / "per_sample_results.csv")
    require(
        samples.groupby("engine").valid.sum().to_dict() == {"V10": 512, "V13": 512},
        "default edge detection completeness",
    )
    psd = pd.read_csv(output / "PSD/per_structure_psd_summary.csv")
    require(
        set(psd.engine) == {"V10", "V13"} and psd.status.eq("OK").all(),
        "both PSD engines should have support",
    )
    require(
        psd.group_size.eq(1).all() and not psd.spectrum_averaging.any(),
        "no PSD averaging",
    )
    require(
        psd.average_range_px.eq(1).all() and psd.smoothing_pixel.eq(1).all(),
        "single-row PSD extraction",
    )
    require(not psd.signal.str.contains("group").any(), "no grouped width signal")
    for engine in ("V10", "V13"):
        require((output / f"PSD/PSD_{engine}.xlsx").is_file(), engine + " PSD workbook")
    cases["default_trench"]["accepted_edges"] = (
        samples.groupby("engine").valid.sum().to_dict()
    )
    cases["default_trench"]["psd_curves_by_engine"] = (
        psd.groupby("engine").size().to_dict()
    )
    cases["default_trench"]["psd_raw_slots"] = sorted(set(psd.used_slots.astype(int)))
    _, line = execute(
        "default_line", pattern="line", extra=["--skip-psd", "--no-annotated-images"]
    )
    for metric in app.METRICS:
        require(
            abs(
                line[f"旋转_mixed_{metric}_nm"].iloc[0]
                - frame[f"旋转_mixed_{metric}_nm"].iloc[0]
            )
            < 1e-9,
            "line/trench polarity equivalence",
        )

    _, legacy = execute(
        "legacy_reproduction",
        extra=[
            "--skip-psd",
            "--no-annotated-images",
            "--no-auto-roi",
            "--end-trim-fraction",
            "0",
            "--locator-mode",
            "dark-line",
            "--threshold-search",
            "legacy",
        ],
    )
    expected = [
        59.819355890345584,
        2.294909515513369,
        2.3071893009682287,
        0.7311518974546287,
    ]
    for metric, value in zip(app.METRICS, expected):
        require(
            abs(legacy[f"旋转_mixed_{metric}_nm"].iloc[0] - value) < 1e-9,
            "legacy compatibility " + metric,
        )
    _, single = execute(
        "one_group",
        extra=["--skip-psd", "--no-annotated-images", "--group-size", "128"],
    )
    require(
        single.status.iloc[0] == "REVIEW"
        and pd.isna(single["旋转_mixed_LWR_nm"].iloc[0]),
        "single group must be REVIEW/NaN",
    )
    _, optional = execute(
        "path_erf_group8",
        pattern="line",
        extra=[
            "--skip-psd",
            "--no-annotated-images",
            "--viterbi",
            "1",
            "--erf-fit",
            "1",
            "--group-size",
            "8",
        ],
    )
    require(optional.status.iloc[0] == "OK", "optional Viterbi/ERF preserved")

    image = app.regions.read_gray_png(BASE / "examples/input_trench/five/test.png")
    padded = np.full((900, 1100), 160, np.uint8)
    padded[70:582, 510:1022] = image
    root = temp / "padded_input"
    app.regions.unicode_imwrite(root / "offset.png", padded)
    output, padded_frame = execute("offset_foreground", root=root, extra=["--skip-psd"])
    require(
        padded_frame.status.iloc[0] == "OK", "off-center foreground should be detected"
    )
    roi = pd.read_csv(output / "roi_summary.csv").iloc[0]
    require(
        roi.roi_x0 >= 500 and roi.roi_y0 >= 65 and roi.roi_y1 <= 587,
        "exclude exterior background",
    )
    require((output / "ROI/offset.png").is_file(), "ROI diagnostic overlay")
    cases["offset_foreground"]["roi_bounds"] = [
        int(roi[k]) for k in ("roi_x0", "roi_x1", "roi_y0", "roi_y1")
    ]

    gapped = image.copy()
    gapped[240:260] = 160
    root = temp / "gap_input"
    app.regions.unicode_imwrite(root / "gap.png", gapped)
    output, _ = execute(
        "background_gap_force_fill",
        root=root,
        extra=["--edge-continuity", "100", "--no-annotated-images"],
    )
    sample = pd.read_csv(output / "per_sample_results.csv")
    excluded = sample.background_excluded.fillna(False)
    require(
        excluded.any() and not sample.loc[excluded, "valid"].any(),
        "continuity must not fill background",
    )
    pitch_audit = pd.read_csv(output / "pitch_samples.csv")
    require(
        pitch_audit.background_excluded.any()
        and not pitch_audit.loc[pitch_audit.background_excluded, "valid"].any()
        and pitch_audit.loc[pitch_audit.background_excluded, "旋转_pitch_CD_nm"]
        .isna()
        .all(),
        "pitch cannot use background even with forced continuity",
    )
    psd_audit = pd.read_csv(output / "PSD/edge_coordinates.csv")
    require(
        not psd_audit.loc[psd_audit.background_excluded, "valid"].any(),
        "PSD must exclude background",
    )

    root = temp / "blank_input"
    app.regions.unicode_imwrite(root / "blank.png", np.full_like(image, 130))
    output, blank = execute(
        "blank_background", root=root, extra=["--no-annotated-images"]
    )
    require(
        blank.status.iloc[0] == "SKIPPED_BACKGROUND"
        and pd.isna(blank["旋转_mixed_CD_nm"].iloc[0]),
        "blank is skipped, not zero",
    )
    require(
        (output / "PSD/PSD_V10.xlsx").is_file(),
        "empty PSD output still has schema/workbook",
    )

    # Orchestration fault tests reuse known real edges; no source mutation.
    for fault in (
        "v10_failure",
        "v13_failure",
        "both_engines_failure",
        "annotation_failure",
        "psd_failure",
        "pitch_failure",
        "unmatched_machine",
        "missing_machine",
    ):

        def reuse(path, params, **kwargs):
            if (
                fault == params.engine_kind.lower() + "_failure"
                or fault == "both_engines_failure"
            ):
                raise RuntimeError("injected engine failure")
            return deepcopy(cached[params.engine_kind])

        extras = [] if fault == "psd_failure" else ["--skip-psd"]
        if fault != "annotation_failure":
            extras += ["--no-annotated-images"]
        if fault in ("unmatched_machine", "missing_machine"):
            machine = temp / (fault + ".xlsx")
            if fault == "unmatched_machine":
                book = Workbook()
                book.active.title = "v2"
                book.active.append(["A", "name", "C", "CD", "LER", "LWR"])
                book.active.append([None, "absent.png", None, 60, 2, 1])
                book.save(machine)
            extras += ["--machine-excel", str(machine)]
        with patch.object(app.edges, "measure_image", side_effect=reuse):
            if fault == "annotation_failure":
                with patch.object(
                    app.edges,
                    "save_annotated_image",
                    side_effect=OSError("injected annotation failure"),
                ):
                    output, result = execute(fault, extra=extras)
            elif fault == "pitch_failure":
                with patch.object(
                    app.pitch,
                    "measure_pitch",
                    side_effect=RuntimeError("injected pitch failure"),
                ):
                    output, result = execute(fault, extra=extras)
                require(
                    result.status.eq("REVIEW").all()
                    and result["旋转_pitch_CD_nm"].isna().all(),
                    "pitch failure retains primary metrics and empty pitch",
                )
            elif fault == "psd_failure":
                with patch.object(
                    app.PSDBatch,
                    "add_measurement",
                    side_effect=RuntimeError("injected PSD failure"),
                ):
                    output, result = execute(fault, extra=extras)
            else:
                output, result = execute(fault, extra=extras)
        require(len(result) == 1, "one row per image after faults")
        if fault == "v10_failure":
            require(
                pd.notna(result["旋转_mixed_CD_nm"].iloc[0])
                and pd.isna(result["旋转_mixed_LER_left_nm"].iloc[0]),
                "retain V13 partial result",
            )
        elif fault == "v13_failure":
            require(
                pd.notna(result["旋转_mixed_LER_left_nm"].iloc[0])
                and pd.isna(result["旋转_mixed_CD_nm"].iloc[0]),
                "retain V10 partial result",
            )
        elif fault == "both_engines_failure":
            require(
                result.status.iloc[0] == "ERROR"
                and pd.isna(result["旋转_mixed_CD_nm"].iloc[0]),
                "all-error schema retained",
            )
        else:
            require(
                pd.notna(result["旋转_mixed_CD_nm"].iloc[0]),
                "auxiliary failure must preserve CD",
            )
        if fault == "unmatched_machine":
            require(
                (output / "machine_unmatched_rows.csv").is_file(),
                "unmatched reference rows must be exported",
            )
    root = temp / "corrupt_input"
    root.mkdir()
    (root / "broken.png").write_bytes(b"not a PNG file")
    _, corrupt = execute("unreadable_png", root=root, extra=["--skip-psd"])
    require(corrupt.status.iloc[0] == "ERROR", "unreadable PNG must appear in summary")

    # A rerun must not overwrite an existing result, even for the same version.
    previous = temp / "default_trench" / "image_summary.csv"
    content = previous.read_bytes()
    with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
        rc = app.main(
            [
                "--root",
                str(BASE / "examples/input_trench"),
                "--output",
                str(previous.parent),
                "--pattern",
                "trench",
            ]
        )
    require(rc == 1 and previous.read_bytes() == content, "nonempty output protection")
    cases["output_protection"] = dict(exit_code=rc, original_csv_preserved=True)
    from check_localization import integration_tests as locator_tests

    cases["v115_localization"] = locator_tests(temp, execute, require)
    from check_line_pitch import integration_tests as pitch_tests

    cases["line_pitch"] = pitch_tests(temp, execute, require)
    return cases


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, np.generic):
        return clean(value.item())
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


if __name__ == "__main__":
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--quick", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = dict(
        version="V1_15",
        python=platform.python_version(),
        dependencies={
            p: importlib.metadata.version(p)
            for p in (
                "numpy",
                "scipy",
                "pandas",
                "opencv-python",
                "openpyxl",
                "matplotlib",
            )
        },
        analytical=analytical_tests(),
    )
    print("Analytical/background tests: PASS", flush=True)
    if not args.quick:
        with tempfile.TemporaryDirectory(prefix="measure-v115-check-") as temp:
            result["integration"] = integration_tests(Path(temp))
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(clean(result), ensure_ascii=False, indent=2, allow_nan=False)
            + "\n"
        )
    print("V1_15 SELF-CHECK PASSED", flush=True)
