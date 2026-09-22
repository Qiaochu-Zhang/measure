#!/usr/bin/env python3
"""Check V1_18 full-population trenches, unchanged line/edge math and standalone ZIP."""

import ast
import hashlib
import json
from pathlib import Path
import shlex
import subprocess
import sys
import tempfile
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "sem_cd_measure_200k_batch_V1_18_complete_package"
BASELINE = ROOT / "sem_cd_measure_200k_batch_V1_17_complete_package"
BASELINE_COMMIT = "2301b355b08765397d386ad33f7412eaa4b376c6"


def run(args, cwd):
    result = subprocess.run(
        [sys.executable, *map(str, args)],
        cwd=cwd,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return result.stdout


def verify_migration():
    files = subprocess.check_output(
        ["git", "ls-tree", "-r", "--name-only", BASELINE_COMMIT], cwd=ROOT, text=True
    ).splitlines()
    preserved = [
        name
        for name in files
        if any(
            v in name.lower()
            for v in ("v1_15", "v1_16", "v1_17", "v115", "v116", "v117")
        )
    ]
    for name in preserved:
        expected = subprocess.check_output(
            ["git", "show", f"{BASELINE_COMMIT}:{name}"], cwd=ROOT
        )
        assert (ROOT / name).read_bytes() == expected, "Old version changed: " + name
    return dict(
        status="PASS",
        baseline_commit=BASELINE_COMMIT,
        old_version_files_verified=len(preserved),
        old_zip_sha256={
            version: hashlib.sha256(
                (
                    ROOT / f"sem_cd_measure_200k_batch_{version}_complete_package.zip"
                ).read_bytes()
            ).hexdigest()
            for version in ("V1_15", "V1_16", "V1_17")
        },
    )


def compare_source_edges(old_output, new_output):
    from check_trench_selection import RAW_FIELDS

    old = pd.read_csv(old_output / "per_sample_results.csv")
    new = pd.read_csv(new_output / "trench_source_samples.csv")
    compared = []
    for (engine, tid), source in new.groupby(["engine", "source_trench_id"]):
        groups = list(old[old.engine == engine].groupby("space_index"))
        if not groups:
            continue
        distances = [
            (
                abs(
                    group["未旋转_left_edge_x_px"].mean()
                    - source["未旋转_left_edge_x_px"].mean()
                ),
                index,
                group,
            )
            for index, group in groups
        ]
        distance, index, group = min(distances, key=lambda item: item[0])
        if distance > 5:
            continue  # Newly scanned trench did not belong to the old capped set.
        pd.testing.assert_frame_equal(
            group[RAW_FIELDS].reset_index(drop=True),
            source[RAW_FIELDS].reset_index(drop=True),
            check_exact=True,
        )
        compared.append(
            dict(
                engine=engine,
                source_trench_id=int(tid),
                old_space_index=int(index),
                samples=len(source),
                qualified=bool(source.source_trench_qualified.iloc[0]),
            )
        )
    assert {r["engine"] for r in compared} == {"V10", "V13"}
    return compared


def verify_compatibility(temp):
    report = {}
    cases = [
        ("line_default", "input_dark_line", "line", 100, []),
        (
            "line_optional",
            "input_dark_line",
            "line",
            100,
            [
                "--viterbi",
                "1",
                "--erf-fit",
                "1",
                "--sample-number",
                "32",
                "--threshold-left",
                "40",
                "--threshold-right",
                "60",
            ],
        ),
        ("trench_healthy", "input_dark_line", "trench", 60, []),
        (
            "trench_legacy_quality",
            "input_trench",
            "trench",
            60,
            [
                "--max-number",
                "4",
                "--no-auto-roi",
                "--end-trim-fraction",
                "0",
                "--locator-mode",
                "dark-line",
                "--threshold-search",
                "legacy",
            ],
        ),
        (
            "trench_false_center",
            "input_dark_line",
            "trench",
            60,
            ["--center-x", "620", "--meas-area-width", "900"],
        ),
    ]
    for name, fixture, pattern, reference, extra in cases:
        print("compatibility:", name, flush=True)
        outputs = []
        for version, package in [("V1_17", BASELINE), ("V1_18", PACKAGE)]:
            output = temp / (name + version)
            run(
                [
                    package / f"sem_cd_measure_200k_batch_{version}.py",
                    "--root",
                    PACKAGE / "examples" / fixture,
                    "--output",
                    output,
                    "--pattern",
                    pattern,
                    "--pixel-size",
                    "1",
                    f"--{pattern}-reference-nm",
                    str(reference),
                    "--skip-psd",
                    "--skip-statistics-plots",
                    "--no-auto-machine-comparison",
                    "--no-annotated-images",
                    *extra,
                ],
                ROOT,
            )
            outputs.append(output)
        old, new = [pd.read_csv(out / "measurement_summary.csv") for out in outputs]
        if pattern == "line":
            tables = {}
            for name_ in [
                "per_sample_results.csv",
                "line_trench_pairs.csv",
                "line_source_samples.csv",
                "pitch_periods.csv",
                "pitch_samples.csv",
            ]:
                before, after = [pd.read_csv(out / name_) for out in outputs]
                pd.testing.assert_frame_equal(
                    before, after[before.columns], check_exact=True
                )
                tables[name_] = len(before)
            columns = [
                c
                for c in old
                if c.startswith(("旋转", "未旋转", "pitch_"))
                or c in ("status", "method")
            ]
            pd.testing.assert_frame_equal(old[columns], new[columns], check_exact=True)
            assert new.status.eq("OK").all()
            report[name] = dict(
                status="PASS", exact_tables=tables, exact_summary_columns=columns
            )
        else:
            sources = compare_source_edges(*outputs)
            selection = pd.read_csv(outputs[1] / "trench_selection.csv")
            assert not selection.loc[~selection.qualified, "selected"].any()
            if name == "trench_healthy":
                columns = [
                    c
                    for c in old
                    if c.startswith(("旋转", "未旋转")) and "pitch" not in c
                ]
                pd.testing.assert_frame_equal(
                    old[columns], new[columns], check_exact=True
                )
                assert new.status.eq("OK").all()
            elif name == "trench_legacy_quality":
                assert not selection.loc[selection.engine == "V13", "qualified"].any()
                assert new.loc[new.method == "V13", "旋转_CD_nm"].isna().all()
                assert new.status.eq("REVIEW").all()
            else:
                before = pd.read_csv(outputs[0] / "locator_candidates.csv")
                center = before[before.selected & before.selected_role.eq("center")]
                assert len(center) == 2 and not center.brightness_true_trench.any()
                assert selection.source_trench_id.mod(2).eq(0).all()
                assert (
                    abs(new.loc[new.method == "mixed", "旋转_CD_nm"].iloc[0] - 60) < 1.2
                )
            report[name] = dict(
                status="PASS",
                exact_common_source_edges=sources,
                old_summary=json.loads(
                    old[["method", "status", "旋转_CD_nm", "旋转_pitch_CD_nm"]].to_json(
                        orient="records"
                    )
                ),
                new_summary=json.loads(
                    new[["method", "status", "旋转_CD_nm", "旋转_pitch_CD_nm"]].to_json(
                        orient="records"
                    )
                ),
            )
    return report


def main():
    sys.path.insert(0, str(PACKAGE))
    import sem_cd_measure_200k_batch_V1_18 as app

    report = {"version_migration": verify_migration()}
    guide = (ROOT / "V1_18_RUN_GUIDE.md").read_text()
    actions = app.build_parser()._actions
    for action in actions:
        assert any(f"`{name}`" in guide for name in action.option_strings), (
            action.option_strings
        )
    examples = []
    for line in guide.splitlines():
        if (
            line.startswith("python sem_cd_measure_200k_batch_V1_18.py ")
            and "--help" not in line
        ):
            args = shlex.split(line)[2:]
            app.validate_args(app.build_parser().parse_args(args), args)
            examples.append(line)
    report["documentation"] = dict(
        cli_options=len(actions), validated_examples=len(examples)
    )
    unchanged = []
    for name in (
        "cdsem_locator.py",
        "cdsem_localization.py",
        "cdsem_lines.py",
        "cdsem_refinement.py",
        "cdsem_statistics.py",
        "cdsem_psd.py",
    ):
        old = (BASELINE / name).read_text().replace("V1_17", "V1_18")
        new = (PACKAGE / name).read_text()
        assert ast.dump(ast.parse(old)) == ast.dump(ast.parse(new)), name
        unchanged.append(name)
    report["unchanged_algorithm_modules_after_version_rename"] = unchanged
    functions = []
    for package in (BASELINE, PACKAGE):
        tree = ast.parse((package / "cdsem_engine.py").read_text())
        functions.append(
            {
                node.name: ast.dump(node)
                for node in tree.body
                if isinstance(node, ast.FunctionDef)
            }
        )
    changed = [
        name for name in functions[0] if functions[0][name] != functions[1].get(name)
    ]
    assert changed == [
        "measure_image",
    ], changed
    report["edge_engine_changed_functions"] = changed
    metric_functions = []
    for package, version in [(BASELINE, "V1_17"), (PACKAGE, "V1_18")]:
        tree = ast.parse(
            (package / f"sem_cd_measure_200k_batch_{version}.py").read_text()
        )
        metric_functions.append(
            {n.name: ast.dump(n) for n in tree.body if isinstance(n, ast.FunctionDef)}
        )
    preserved_metrics = [
        name
        for name in metric_functions[0]
        if name.startswith(
            ("coordinate_metrics", "aggregate", "choose_aggregate", "fit_common")
        )
    ]
    assert preserved_metrics
    for name in preserved_metrics:
        assert metric_functions[0][name] == metric_functions[1][name], name
    report["unchanged_main_statistics_functions"] = preserved_metrics
    archive = ROOT / (PACKAGE.name + ".zip")
    with tempfile.TemporaryDirectory(prefix="v118-release-") as temp:
        temp = Path(temp)
        report["numerical_compatibility"] = verify_compatibility(temp)
        print("archive: standalone extraction and exports", flush=True)
        with zipfile.ZipFile(archive) as zipped:
            assert zipped.testzip() is None
            zipped.extractall(temp)
        unpacked = temp / PACKAGE.name
        manifest = (unpacked / "PACKAGE_SHA256.txt").read_text().splitlines()
        for line in manifest:
            digest, path = line.split("  ", 1)
            assert (
                hashlib.sha256((unpacked / path).read_bytes()).hexdigest() == digest
            ), path
        report["archive"] = dict(
            files_verified=len(manifest),
            sha256=hashlib.sha256(archive.read_bytes()).hexdigest(),
        )
        report["unpacked_quick_self_check"] = run(
            ["self_check_V1_18.py", "--quick"], unpacked
        ).strip()
        book = Workbook()
        book.active.title = "v2"
        book.active.append(["A", "name", "C", "CD", "LER", "LWR"])
        book.active.append([None, "test.png", None, 60, 2, 1])
        book.save(unpacked / "reference.xlsx")
        output = unpacked / "full_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_18.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                output,
                "--pattern",
                "trench",
                "--pixel-size",
                "1",
                "--trench-reference-nm",
                "60",
                "--machine-excel",
                "reference.xlsx",
                "--save-debug-masks",
                "--viterbi",
                "1",
                "--erf-fit",
                "1",
            ],
            unpacked,
        )
        settings = json.loads((output / "settings.json").read_text())
        assert (
            settings["script_version"] == "V1_18"
            and settings["patch_version"] == "V1_18"
        )
        assert not list(output.glob("*V1_15*")), "Old release name in V1_18 output"
        images = pd.read_csv(output / "image_summary.csv")
        errors = pd.read_csv(output / "processing_errors.csv")
        machine = pd.read_csv(output / "machine_comparison_detailed.csv")
        assert images.status.eq("OK").all() and errors.empty
        assert images.locator_mode_selected.eq("dark-line").all()
        assert len(machine) > 0
        psd = pd.read_csv(output / "PSD/per_structure_psd_summary.csv")
        assert set(psd.engine) == {"V10", "V13"}
        assert psd.status.eq("OK").all()
        png_count = len(list(output.rglob("*.png")))
        assert png_count > 10
        report["unpacked_full_export"] = dict(
            status="PASS",
            viterbi=True,
            erf=True,
            image_statuses=list(images.status),
            locator_modes=list(images.locator_mode_selected),
            processing_errors=len(errors),
            machine_rows=len(machine),
            psd_curves=len(psd),
            png_exports=png_count,
            mixed_cd_nm=float(images["旋转_mixed_CD_nm"].iloc[0]),
        )
        from check_trench_selection import verify_trench_exports

        def require(condition, reason):
            assert condition, reason

        report["unpacked_full_export"]["qualified_trenches"] = verify_trench_exports(
            output, require, 3
        )
        summary = pd.read_csv(output / "measurement_summary.csv")
        assert summary.columns[-1] == "旋转_pitch_CD_nm"
        assert summary.pitch_count.eq(3).all()
        assert (summary["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        report["unpacked_full_export"]["pitch_count"] = list(summary.pitch_count)
        report["unpacked_full_export"]["rotated_pitch_cd_nm"] = list(
            summary["旋转_pitch_CD_nm"]
        )
        line_output = unpacked / "line_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_18.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                line_output,
                "--pattern",
                "line",
                "--pixel-size",
                "1",
                "--line-reference-nm",
                "100",
                "--no-auto-machine-comparison",
                "--skip-statistics-plots",
            ],
            unpacked,
        )
        line = pd.read_csv(line_output / "measurement_summary.csv")
        errors = pd.read_csv(line_output / "processing_errors.csv")
        psd = pd.read_csv(line_output / "PSD/per_structure_psd_summary.csv")
        assert line.status.eq("OK").all() and errors.empty
        assert abs(line.loc[line.method == "V13", "旋转_CD_nm"].iloc[0] - 100) < 1
        assert (
            line.pitch_count.eq(3).all()
            and (line["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        )
        assert set(psd.engine) == {"V10", "V13"} and psd.status.eq("OK").all()
        book = load_workbook(
            line_output / "CD_measurement_200K_V1_18_results.xlsx", read_only=True
        )
        assert next(book["measurement_summary"].values)[-1] == "旋转_pitch_CD_nm"
        assert (
            "line_trench_pairs" in book.sheetnames
            and "line_source_samples" in book.sheetnames
        )
        book.close()
        from check_shared_lines import verify_shared_exports

        def require(condition, reason):
            assert condition, reason

        shared = verify_shared_exports(line_output, require)
        report["unpacked_stained_line_export"] = dict(
            status="PASS",
            shared_edges=shared,
            viterbi=False,
            erf=False,
            processing_errors=len(errors),
            psd_curves=len(psd),
            rotated_pitch_cd_nm=list(line["旋转_pitch_CD_nm"]),
            pitch_count=list(line.pitch_count),
            mixed_cd_nm=float(line.loc[line.method == "mixed", "旋转_CD_nm"].iloc[0]),
        )
    evidence = ROOT / "audits/v1_18_release_evidence.json"
    evidence.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
