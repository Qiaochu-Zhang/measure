#!/usr/bin/env python3
"""Check V1_16 migration, original V1_15 preservation, docs and the unpacked ZIP."""

import ast
import hashlib
import json
from pathlib import Path
import shlex
import subprocess
import sys
import tempfile
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "sem_cd_measure_200k_batch_V1_16_complete_package"
BASELINE = ROOT / "sem_cd_measure_200k_batch_V1_15_complete_package"
ORIGINAL_COMMIT = "41e3294092016a8e6ca44c4d3a6217bda2e6ea90"
PATCH_COMMIT = "ae11296b1e09857bf97558a0c395ad5e8fadec06"


def run(args, cwd):
    result = subprocess.run(
        [sys.executable, *map(str, args)],
        cwd=cwd,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return result.stdout


def verify_migration():
    old_paths = [
        BASELINE.name,
        BASELINE.name + ".zip",
        "V1_15_RELEASE_NOTES.md",
        "V1_15_RUN_GUIDE.md",
        "audits/package_v115.py",
        "audits/validate_v115_release.py",
        "audits/v1_15_release_evidence.json",
    ]
    files = subprocess.check_output(
        ["git", "ls-tree", "-r", "--name-only", ORIGINAL_COMMIT, "--", *old_paths],
        cwd=ROOT,
        text=True,
    ).splitlines()
    for name in files:
        expected = subprocess.check_output(
            ["git", "show", f"{ORIGINAL_COMMIT}:{name}"], cwd=ROOT
        )
        assert (ROOT / name).read_bytes() == expected, "Original V1_15 changed: " + name
    assert not subprocess.check_output(
        ["git", "diff", "--name-only", ORIGINAL_COMMIT, "--", *old_paths],
        cwd=ROOT,
        text=True,
    ).strip()
    modules = []
    for target in sorted(PACKAGE.glob("*.py")):
        if target.name.startswith(("self_check_", "check_")):
            continue
        name = target.name.replace("V1_16", "V1_15")
        old = subprocess.check_output(
            ["git", "show", f"{PATCH_COMMIT}:{BASELINE.name}/{name}"],
            cwd=ROOT,
            text=True,
        )
        normalized = old.replace("V1_15", "V1_16").replace(
            "V1_16_line_pitch_1", "V1_16"
        )
        normalized = normalized.replace(
            "metadata_first_methods_raw_psd_v115_pitch",
            "metadata_first_methods_raw_psd_v116_pitch",
        )
        if name == "cdsem_regions.py":
            normalized = normalized.replace(
                '                    "V1_16",',
                '                    "V1_15",\n                    "V1_16",',
            )
        assert ast.dump(ast.parse(normalized)) == ast.dump(
            ast.parse(target.read_text())
        ), target.name
        modules.append(target.name)
    return {
        "original_v115_commit": ORIGINAL_COMMIT,
        "original_v115_files_verified": len(files),
        "original_v115_zip_sha256": hashlib.sha256(
            (ROOT / (BASELINE.name + ".zip")).read_bytes()
        ).hexdigest(),
        "feature_source_commit": PATCH_COMMIT,
        "preserved_algorithm_modules": modules,
        "allowed_migration_changes": [
            "V1_16 release names and output schema",
            "recognize both V1_15 and V1_16 output directories",
        ],
        "status": "PASS",
    }


def main():
    sys.path.insert(0, str(PACKAGE))
    import sem_cd_measure_200k_batch_V1_16 as app

    report = {"version_migration": verify_migration()}
    guide = (ROOT / "V1_16_RUN_GUIDE.md").read_text()
    actions = app.build_parser()._actions
    for action in actions:
        assert any(f"`{name}`" in guide for name in action.option_strings), (
            action.option_strings
        )
    examples = []
    for line in guide.splitlines():
        if (
            line.startswith("python sem_cd_measure_200k_batch_V1_16.py ")
            and "--help" not in line
        ):
            args = shlex.split(line)[2:]
            app.validate_args(app.build_parser().parse_args(args), args)
            examples.append(line)
    report["documentation"] = dict(
        cli_options=len(actions), validated_examples=len(examples)
    )
    unchanged = []
    for name in (
        "cdsem_locator.py",
        "cdsem_refinement.py",
        "cdsem_statistics.py",
        "cdsem_psd.py",
    ):
        old = (BASELINE / name).read_text().replace("V1_15", "V1_16")
        new = (PACKAGE / name).read_text()
        assert ast.dump(ast.parse(old)) == ast.dump(ast.parse(new)), name
        unchanged.append(name)
    report["unchanged_algorithm_modules_after_version_rename"] = unchanged
    functions = []
    for package in (BASELINE, PACKAGE):
        tree = ast.parse((package / "cdsem_engine.py").read_text())
        functions.append(
            {
                node.name: ast.dump(node)
                for node in tree.body
                if isinstance(node, ast.FunctionDef)
            }
        )
    changed = [
        name for name in functions[0] if functions[0][name] != functions[1].get(name)
    ]
    assert changed == [
        "detect_edge_pair",
        "build_v115_fixed",
        "detect_space_candidates",
    ], changed
    report["edge_engine_changed_functions"] = changed
    archive = ROOT / (PACKAGE.name + ".zip")
    with tempfile.TemporaryDirectory(prefix="v116-release-") as temp:
        temp = Path(temp)
        with zipfile.ZipFile(archive) as zipped:
            assert zipped.testzip() is None
            zipped.extractall(temp)
        unpacked = temp / PACKAGE.name
        manifest = (unpacked / "PACKAGE_SHA256.txt").read_text().splitlines()
        for line in manifest:
            digest, path = line.split("  ", 1)
            assert (
                hashlib.sha256((unpacked / path).read_bytes()).hexdigest() == digest
            ), path
        report["archive"] = dict(
            files_verified=len(manifest),
            sha256=hashlib.sha256(archive.read_bytes()).hexdigest(),
        )
        report["unpacked_quick_self_check"] = run(
            ["self_check_V1_16.py", "--quick"], unpacked
        ).strip()
        book = Workbook()
        book.active.title = "v2"
        book.active.append(["A", "name", "C", "CD", "LER", "LWR"])
        book.active.append([None, "test.png", None, 60, 2, 1])
        book.save(unpacked / "reference.xlsx")
        output = unpacked / "full_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_16.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                output,
                "--pattern",
                "trench",
                "--pixel-size",
                "1",
                "--trench-reference-nm",
                "60",
                "--machine-excel",
                "reference.xlsx",
                "--save-debug-masks",
            ],
            unpacked,
        )
        settings = json.loads((output / "settings.json").read_text())
        assert (
            settings["script_version"] == "V1_16"
            and settings["patch_version"] == "V1_16"
        )
        assert not list(output.glob("*V1_15*")), "Old release name in V1_16 output"
        images = pd.read_csv(output / "image_summary.csv")
        errors = pd.read_csv(output / "processing_errors.csv")
        machine = pd.read_csv(output / "machine_comparison_detailed.csv")
        assert images.status.eq("OK").all() and errors.empty
        assert images.locator_mode_selected.eq("dark-line").all()
        assert len(machine) > 0
        psd = pd.read_csv(output / "PSD/per_structure_psd_summary.csv")
        assert set(psd.engine) == {"V10", "V13"}
        assert psd.status.eq("OK").all()
        png_count = len(list(output.rglob("*.png")))
        assert png_count > 10
        report["unpacked_full_export"] = dict(
            status="PASS",
            image_statuses=list(images.status),
            locator_modes=list(images.locator_mode_selected),
            processing_errors=len(errors),
            machine_rows=len(machine),
            psd_curves=len(psd),
            png_exports=png_count,
            mixed_cd_nm=float(images["旋转_mixed_CD_nm"].iloc[0]),
        )
        summary = pd.read_csv(output / "measurement_summary.csv")
        assert summary.columns[-1] == "旋转_pitch_CD_nm"
        assert summary.pitch_count.eq(3).all()
        assert (summary["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        report["unpacked_full_export"]["pitch_count"] = list(summary.pitch_count)
        report["unpacked_full_export"]["rotated_pitch_cd_nm"] = list(
            summary["旋转_pitch_CD_nm"]
        )
        line_output = unpacked / "line_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_16.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                line_output,
                "--pattern",
                "line",
                "--pixel-size",
                "1",
                "--line-reference-nm",
                "100",
                "--no-auto-machine-comparison",
                "--skip-statistics-plots",
                "--viterbi",
                "1",
                "--erf-fit",
                "1",
            ],
            unpacked,
        )
        line = pd.read_csv(line_output / "measurement_summary.csv")
        errors = pd.read_csv(line_output / "processing_errors.csv")
        psd = pd.read_csv(line_output / "PSD/per_structure_psd_summary.csv")
        assert line.status.eq("OK").all() and errors.empty
        assert abs(line.loc[line.method == "V13", "旋转_CD_nm"].iloc[0] - 100) < 1
        assert (
            line.pitch_count.eq(3).all()
            and (line["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        )
        assert set(psd.engine) == {"V10", "V13"} and psd.status.eq("OK").all()
        book = load_workbook(
            line_output / "CD_measurement_200K_V1_16_results.xlsx", read_only=True
        )
        assert next(book["measurement_summary"].values)[-1] == "旋转_pitch_CD_nm"
        book.close()
        report["unpacked_stained_line_export"] = dict(
            status="PASS",
            viterbi=True,
            erf=True,
            processing_errors=len(errors),
            psd_curves=len(psd),
            rotated_pitch_cd_nm=list(line["旋转_pitch_CD_nm"]),
            pitch_count=list(line.pitch_count),
            mixed_cd_nm=float(line.loc[line.method == "mixed", "旋转_CD_nm"].iloc[0]),
        )
    evidence = ROOT / "audits/v1_16_release_evidence.json"
    evidence.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
