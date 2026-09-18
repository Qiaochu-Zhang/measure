#!/usr/bin/env python3
"""Audit original V1_13 versus V13_modified in separate Python processes.

Run twice with --package original/modified, then use --combine OLD.json NEW.json.
Fault injection reuses real baseline edge measurements; it tests orchestration,
not SEM accuracy. Source files and original inputs are never changed.
"""
from __future__ import annotations

import argparse
import ast
from collections import Counter
from contextlib import ExitStack, redirect_stderr, redirect_stdout
from copy import deepcopy
from io import StringIO
import hashlib
import importlib
import importlib.metadata
import json
from pathlib import Path
import platform
import subprocess
import sys
import tempfile
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
PACKAGES = {
    "original": ROOT / "sem_cd_measure_200k_batch_V1_13_complete_package",
    "modified": ROOT / "sem_cd_measure_200k_batch_V13_modified_complete_package",
}
ENTRIES = {
    "original": "sem_cd_measure_200k_batch_V1_13",
    "modified": "sem_cd_measure_200k_batch_V13_modified",
}
METRICS = ("CD", "LER_left", "LER_right", "LWR")


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [clean(v) for v in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return clean(value.item())
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def read_csv(path):
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def summarize(output, package, rc, log):
    frame = read_csv(output / "image_summary.csv")
    result = dict(exit_code=rc, image_summary_exists=(output / "image_summary.csv").exists(),
                  image_summary_rows=len(frame), columns=list(frame.columns), images=[])
    for _, row in frame.iterrows():
        metrics = {}
        for mode, chinese in [("rotated", "旋转"), ("unrotated", "未旋转")]:
            for metric in METRICS:
                column = (f"{mode}_{metric}_nm" if package == "original"
                          else f"{chinese}_mixed_{metric}_nm")
                metrics[f"{mode}_{metric}_nm"] = row.get(column)
        result["images"].append(dict(image_key=row.get("image_key"), status=row.get("status"),
            metrics=metrics, warning=row.get("warning"), error_message=row.get("error_message")))
    errors = read_csv(output / "processing_errors.csv")
    result["errors"] = errors[[c for c in ("stage", "error_type", "error_message") if c in errors]].to_dict("records")
    result["sample_rows"] = len(read_csv(output / "per_sample_results.csv"))
    objects = read_csv(output / "trench_objects.csv")
    result["object_rows"] = len(objects)
    result["object_selection_flags"] = objects[[c for c in (
        "space_role", "used_for_statistics", "V10_used_for_image_statistics",
        "V13_used_for_image_statistics") if c in objects]].to_dict("records")
    result["coordinate_rows"] = len(read_csv(output / ("coordinate_system_results.csv" if package == "original" else "coordinate_results.csv")))
    settings = output / "settings.json"
    if settings.exists():
        data = json.loads(settings.read_text())
        result["counts"] = {k: data.get(k) for k in ("processed_png_count", "ok_count", "review_count", "error_count", "diagnostic_error_count")}
    spectra = read_csv(output / "PSD/per_trench_psd_summary.csv")
    result["psd_statuses"] = ({f"{engine}:{status}": int(len(g))
        for (engine, status), g in spectra.groupby(["engine", "status"])} if not spectra.empty else {})
    excel = output / ("CD_measurement_200K_V1_13_results.xlsx" if package == "original" else "CD_measurement_200K_V13_modified_results.xlsx")
    result["main_workbook_exists"] = excel.exists()
    if excel.exists():
        book = load_workbook(excel, read_only=True)
        result["sheets"] = book.sheetnames
        book.close()
    result["log_tail"] = log.splitlines()[-5:]
    return result


def run_worker(package, destination):
    directory = PACKAGES[package]
    if not directory.is_dir():
        raise RuntimeError(f"Extract {directory.name}.zip first")
    sys.path.insert(0, str(directory))
    app = importlib.import_module(ENTRIES[package])
    snapshot = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in directory.glob("*.py")}
    parser = app.build_parser()
    evidence = dict(package=package, environment=dict(python=platform.python_version(),
        packages={n: importlib.metadata.version(n) for n in ("numpy", "pandas", "scipy", "matplotlib", "opencv-python", "openpyxl")}),
        source_sha256=snapshot, cli={a.dest: dict(default=a.default, choices=a.choices)
                                    for a in parser._actions if a.dest != "help"}, cases={})
    checks = (["self_check_V1_13.py"] if package == "original"
              else ["self_check_V13_modified.py", "self_check_partial_results.py"])
    evidence["self_checks"] = {}
    for check in checks:
        run = subprocess.run([sys.executable, str(directory/check)], cwd=directory,
                             text=True, capture_output=True, check=True)
        evidence["self_checks"][check] = run.stdout.strip()

    cached = {}
    with tempfile.TemporaryDirectory(prefix=f"measure-{package}-comparison-") as temp:
        temp = Path(temp)

        def argv(output, kind="trench"):
            fixtures = "validation" if package == "original" else "examples"
            return ["--root", str(directory/fixtures/f"input_{kind}"), "--output", str(output),
                    "--pattern", kind, "--pixel-size", "1", f"--{kind}-reference-nm", "60",
                    "--max-number", "4", "--no-auto-machine-comparison", "--skip-statistics-plots"]

        def invoke(name, arguments):
            print(f"{package}: {name}", flush=True)
            log = StringIO()
            with redirect_stdout(log), redirect_stderr(log):
                rc = app.main(arguments)
            output = Path(arguments[arguments.index("--output")+1])
            result = summarize(output, package, rc, log.getvalue())
            evidence["cases"][name] = result
            return result

        # Genuine end-to-end baselines capture real measurements for fault tests.
        for kind in ("trench", "line"):
            with ExitStack() as stack:
                for tag, module in (("V10", app.v5), ("V13", app.v13v5)):
                    real_measure = module.measure_image

                    def capture(path, params, real=real_measure, engine=tag):
                        measured = real(path, params)
                        if kind == "trench":
                            cached[engine] = deepcopy(measured)
                        return measured

                    stack.enter_context(patch.object(module, "measure_image", side_effect=capture))
                invoke(f"baseline_{kind}", argv(temp/f"baseline_{kind}", kind))
        evidence["baseline_edges"] = {}
        for engine, measurement in cached.items():
            evidence["baseline_edges"][engine] = dict(slots=len(measurement.sample_rows),
                accepted=sum(bool(r["valid"]) for r in measurement.sample_rows),
                failures=dict(Counter(r["failure_reason"] for r in measurement.sample_rows if not r["valid"])))

        invoke("one_group", argv(temp/"one_group") + ["--group-size", "128", "--edge-continuity", "20", "--no-annotated-images"])
        invoke("line_path_erf", argv(temp/"line_path_erf", "line") +
               ["--viterbi", "1", "--erf-fit", "1", "--edge-continuity", "50", "--group-size", "8", "--no-annotated-images"])

        cases = ("v10_exception", "v13_exception", "v10_no_candidates", "both_no_candidates",
                 "roles_disjoint", "unselected_object_flag", "unused_v10_group_row_missing", "annotation_failure",
                 "psd_failure", "machine_unmatched", "output_schema_collision")
        for case in cases:
            output = temp/case
            arguments = argv(output)
            if case != "psd_failure":
                arguments += ["--skip-psd"]
            if case != "annotation_failure":
                arguments += ["--no-annotated-images"]
            measurements = deepcopy(cached)
            if case in ("v10_no_candidates", "both_no_candidates"):
                for tag in (("V10", "V13") if case == "both_no_candidates" else ("V10",)):
                    m = measurements[tag]
                    m.space_rows = []
                    m.sample_rows = []
                    m.annotation_payload["spaces"] = []
                    m.image_row.update(valid=False, triplet_complete=False, selected_space_count=0,
                                       stable_space_count=0, warning="injected no candidates")
            if case == "roles_disjoint":
                for tag, m in measurements.items():
                    for row in m.space_rows + m.sample_rows:
                        row["space_role"] = tag + "_" + row["space_role"]
                    for ann in m.annotation_payload["spaces"]:
                        ann["candidate"].role = tag + "_" + ann["candidate"].role
            if case == "unselected_object_flag":
                # Keep all measured coordinates; select three of four objects.
                for m in measurements.values():
                    for i, row in enumerate(m.space_rows):
                        row["stable"] = i != 0
            if case == "machine_unmatched":
                book = Workbook()
                sheet = book.active
                sheet.title = "v2"
                sheet.append(["A", "file", "C", "CD", "LER", "LWR"])
                sheet.append([None, "absent.png", None, 60, 2, 1])
                machine = temp/"unmatched.xlsx"
                book.save(machine)
                arguments += ["--machine-excel", str(machine)]
            if case == "output_schema_collision":
                output.mkdir()
                (output/"image_summary.csv").write_text("previous_result\n123\n")
                (output/"settings.json").write_text('{"script_version":"unrelated_version"}')
            with ExitStack() as stack:
                for tag, module in (("V10", app.v5), ("V13", app.v13v5)):
                    if case == tag.lower() + "_exception":
                        replacement = patch.object(module, "measure_image", side_effect=RuntimeError(f"injected {tag} failure"))
                    else:
                        replacement = patch.object(module, "measure_image", return_value=measurements[tag])
                    stack.enter_context(replacement)
                if case == "annotation_failure":
                    stack.enter_context(patch.object(app.v13v5, "save_annotated_image", side_effect=OSError("injected annotation failure")))
                if case == "psd_failure":
                    real_add = app.PSDBatch.add_measurement

                    def fail_psd(self, measurement, key, engine, params):
                        if engine == "V13":
                            raise RuntimeError("injected V13 PSD failure")
                        return real_add(self, measurement, key, engine, params)

                    stack.enter_context(patch.object(app.PSDBatch, "add_measurement", new=fail_psd))
                if case == "unused_v10_group_row_missing":
                    real_variants = app.v7.process_measurement_variants

                    def drop_unused(measurement, params, *rest):
                        objects, images, samples, base = real_variants(measurement, params, *rest)
                        if type(params) is app.v5.MeasurementParams:
                            images = [r for r in images if r["result_id"] != app.v10.GROUP4_RESULT_ID]
                        return objects, images, samples, base

                    stack.enter_context(patch.object(app.v7, "process_measurement_variants", side_effect=drop_unused))
                result = invoke(case, arguments)
                if case == "unselected_object_flag":
                    result["expected_selected_indices"] = {
                        tag: app.v7.choose_aggregate_space_indices(m.space_rows, 3)
                        for tag, m in measurements.items()}
                    result["excluded_roles"] = {
                        tag: m.space_rows[0]["space_role"] for tag, m in measurements.items()}
                if case == "output_schema_collision":
                    result["previous_summary_preserved"] = (output/"image_summary.csv").read_text() == "previous_result\n123\n"

        if package == "original":
            print("original: via smoke test", flush=True)
            output = temp/"via"
            log = StringIO()
            with redirect_stdout(log), redirect_stderr(log):
                rc = app.main(["--root", str(directory/"validation/input_via"), "--output", str(output),
                               "--pattern", "via", "--pixel-size", "1", "--via-reference-nm", "60"])
            frame = read_csv(output/"image_summary.csv")
            evidence["via"] = dict(exit_code=rc, rows=frame.to_dict("records"),
                                   workbooks=[p.name for p in output.glob("*.xlsx")],
                                   log_tail=log.getvalue().splitlines()[-8:])

    assert all(hashlib.sha256((directory/name).read_bytes()).hexdigest() == digest
               for name, digest in snapshot.items())
    evidence["source_files_unchanged"] = True
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(clean(evidence), ensure_ascii=False, indent=2, allow_nan=False)+"\n")


def combine(paths, destination):
    packages = {data["package"]: data for data in (json.loads(p.read_text()) for p in paths)}
    old, new = packages["original"], packages["modified"]
    common = sorted(set(old["source_sha256"]) & set(new["source_sha256"]))
    identical = [name for name in common if old["source_sha256"][name] == new["source_sha256"][name]]
    psd_trees = [ast.parse((PACKAGES[p]/name).read_text()) for p, name in
                 (("original", "cdsem_v112_psd.py"), ("modified", "cdsem_v13_modified_psd.py"))]
    functions = [{n.name: ast.dump(n, include_attributes=False) for n in tree.body
                  if isinstance(n, ast.FunctionDef)} for tree in psd_trees]
    equal_psd = {k: functions[0][k] == functions[1].get(k) for k in functions[0]}
    differences = {}
    matched_counts = {}
    for case in ("baseline_trench", "baseline_line", "one_group", "line_path_erf"):
        a = old["cases"][case]["images"][0]["metrics"]
        b = new["cases"][case]["images"][0]["metrics"]
        matched = [k for k in a if a[k] is not None and b[k] is not None]
        matched_counts[case] = len(matched)
        differences[case] = max((abs(a[k]-b[k]) for k in matched), default=None)
    fixture_hashes = {}
    for kind in ("trench", "line"):
        hashes = {package: hashlib.sha256((PACKAGES[package]/subdir/f"input_{kind}"/"five/test.png").read_bytes()).hexdigest()
                  for package, subdir in (("original", "validation"), ("modified", "examples"))}
        fixture_hashes[kind] = dict(sha256=hashes, identical=len(set(hashes.values())) == 1)
    evidence = dict(baseline_commit="e51945d", identical_common_python_files=identical,
                    different_common_python_files=sorted(set(common)-set(identical)),
                    fixture_hashes=fixture_hashes,
                    psd_top_level_function_ast_equal=equal_psd,
                    max_abs_mixed_metric_difference_nm=differences,
                    compared_mixed_metric_count=matched_counts, packages=packages)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(clean(evidence), ensure_ascii=False, indent=2, allow_nan=False)+"\n")
    print(f"Wrote {destination}; {len(identical)} identical shared Python files; mixed differences: {differences}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--package", choices=tuple(PACKAGES))
    mode.add_argument("--combine", nargs=2, type=Path, metavar=("OLD_JSON", "NEW_JSON"))
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.package:
        run_worker(args.package, args.output)
    else:
        combine(args.combine, args.output)
