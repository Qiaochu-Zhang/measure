#!/usr/bin/env python3
"""Check V1_17 shared lines, old version preservation, trench parity and ZIP."""

import ast
import hashlib
import json
from pathlib import Path
import shlex
import subprocess
import sys
import tempfile
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "sem_cd_measure_200k_batch_V1_17_complete_package"
BASELINE = ROOT / "sem_cd_measure_200k_batch_V1_16_complete_package"
BASELINE_COMMIT = "c0ae220406543f8a49e9fb880ef9398021ab806f"


def run(args, cwd):
    result = subprocess.run(
        [sys.executable, *map(str, args)],
        cwd=cwd,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return result.stdout


def verify_migration():
    files = subprocess.check_output(
        ["git", "ls-tree", "-r", "--name-only", BASELINE_COMMIT], cwd=ROOT, text=True
    ).splitlines()
    preserved = [
        name
        for name in files
        if any(v in name.lower() for v in ("v1_15", "v1_16", "v115", "v116"))
    ]
    for name in preserved:
        expected = subprocess.check_output(
            ["git", "show", f"{BASELINE_COMMIT}:{name}"], cwd=ROOT
        )
        assert (ROOT / name).read_bytes() == expected, "Old version changed: " + name
    return dict(
        status="PASS",
        baseline_commit=BASELINE_COMMIT,
        old_version_files_verified=len(preserved),
        old_zip_sha256={
            version: hashlib.sha256(
                (
                    ROOT / f"sem_cd_measure_200k_batch_{version}_complete_package.zip"
                ).read_bytes()
            ).hexdigest()
            for version in ("V1_15", "V1_16")
        },
    )


def verify_trench_compatibility(temp):
    report = {}
    for name, fixture, extra in [
        ("stained_default", "input_dark_line", []),
        (
            "bright_optional",
            "input_bright_line",
            ["--viterbi", "1", "--erf-fit", "1", "--group-size", "8"],
        ),
    ]:
        outputs = []
        for version, package in [("V1_16", BASELINE), ("V1_17", PACKAGE)]:
            output = temp / (name + version)
            run(
                [
                    package / f"sem_cd_measure_200k_batch_{version}.py",
                    "--root",
                    PACKAGE / "examples" / fixture,
                    "--output",
                    output,
                    "--pattern",
                    "trench",
                    "--pixel-size",
                    "1",
                    "--trench-reference-nm",
                    "60",
                    "--skip-psd",
                    "--skip-statistics-plots",
                    "--no-auto-machine-comparison",
                    "--no-annotated-images",
                    *extra,
                ],
                ROOT,
            )
            outputs.append(output)
        comparisons = {}
        for filename in [
            "per_sample_results.csv",
            "pitch_periods.csv",
            "pitch_samples.csv",
        ]:
            old, new = [pd.read_csv(out / filename) for out in outputs]
            pd.testing.assert_frame_equal(old, new[old.columns], check_exact=True)
            comparisons[filename] = len(old)
        old, new = [pd.read_csv(out / "measurement_summary.csv") for out in outputs]
        columns = [
            c
            for c in old
            if c.startswith(("旋转", "未旋转", "pitch_")) or c in ("status", "method")
        ]
        pd.testing.assert_frame_equal(old[columns], new[columns], check_exact=True)
        assert old.status.eq("OK").all()
        report[name] = dict(
            status="PASS", exact_tables=comparisons, exact_summary_columns=columns
        )
    return report


def main():
    sys.path.insert(0, str(PACKAGE))
    import sem_cd_measure_200k_batch_V1_17 as app

    report = {"version_migration": verify_migration()}
    guide = (ROOT / "V1_17_RUN_GUIDE.md").read_text()
    actions = app.build_parser()._actions
    for action in actions:
        assert any(f"`{name}`" in guide for name in action.option_strings), (
            action.option_strings
        )
    examples = []
    for line in guide.splitlines():
        if (
            line.startswith("python sem_cd_measure_200k_batch_V1_17.py ")
            and "--help" not in line
        ):
            args = shlex.split(line)[2:]
            app.validate_args(app.build_parser().parse_args(args), args)
            examples.append(line)
    report["documentation"] = dict(
        cli_options=len(actions), validated_examples=len(examples)
    )
    unchanged = []
    for name in (
        "cdsem_locator.py",
        "cdsem_localization.py",
        "cdsem_refinement.py",
        "cdsem_statistics.py",
        "cdsem_psd.py",
    ):
        old = (BASELINE / name).read_text().replace("V1_16", "V1_17")
        new = (PACKAGE / name).read_text()
        assert ast.dump(ast.parse(old)) == ast.dump(ast.parse(new)), name
        unchanged.append(name)
    report["unchanged_algorithm_modules_after_version_rename"] = unchanged
    functions = []
    for package in (BASELINE, PACKAGE):
        tree = ast.parse((package / "cdsem_engine.py").read_text())
        functions.append(
            {
                node.name: ast.dump(node)
                for node in tree.body
                if isinstance(node, ast.FunctionDef)
            }
        )
    changed = [
        name for name in functions[0] if functions[0][name] != functions[1].get(name)
    ]
    assert changed == [
        "detect_space_candidates",
        "measure_image",
    ], changed
    report["edge_engine_changed_functions"] = changed
    metric_functions = []
    for package, version in [(BASELINE, "V1_16"), (PACKAGE, "V1_17")]:
        tree = ast.parse(
            (package / f"sem_cd_measure_200k_batch_{version}.py").read_text()
        )
        metric_functions.append(
            {n.name: ast.dump(n) for n in tree.body if isinstance(n, ast.FunctionDef)}
        )
    preserved_metrics = [
        name
        for name in metric_functions[0]
        if name.startswith(
            ("coordinate_metrics", "aggregate", "choose_aggregate", "fit_common")
        )
    ]
    assert preserved_metrics
    for name in preserved_metrics:
        assert metric_functions[0][name] == metric_functions[1][name], name
    report["unchanged_main_statistics_functions"] = preserved_metrics
    archive = ROOT / (PACKAGE.name + ".zip")
    with tempfile.TemporaryDirectory(prefix="v117-release-") as temp:
        temp = Path(temp)
        report["trench_numerical_compatibility"] = verify_trench_compatibility(temp)
        with zipfile.ZipFile(archive) as zipped:
            assert zipped.testzip() is None
            zipped.extractall(temp)
        unpacked = temp / PACKAGE.name
        manifest = (unpacked / "PACKAGE_SHA256.txt").read_text().splitlines()
        for line in manifest:
            digest, path = line.split("  ", 1)
            assert (
                hashlib.sha256((unpacked / path).read_bytes()).hexdigest() == digest
            ), path
        report["archive"] = dict(
            files_verified=len(manifest),
            sha256=hashlib.sha256(archive.read_bytes()).hexdigest(),
        )
        report["unpacked_quick_self_check"] = run(
            ["self_check_V1_17.py", "--quick"], unpacked
        ).strip()
        book = Workbook()
        book.active.title = "v2"
        book.active.append(["A", "name", "C", "CD", "LER", "LWR"])
        book.active.append([None, "test.png", None, 60, 2, 1])
        book.save(unpacked / "reference.xlsx")
        output = unpacked / "full_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_17.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                output,
                "--pattern",
                "trench",
                "--pixel-size",
                "1",
                "--trench-reference-nm",
                "60",
                "--machine-excel",
                "reference.xlsx",
                "--save-debug-masks",
            ],
            unpacked,
        )
        settings = json.loads((output / "settings.json").read_text())
        assert (
            settings["script_version"] == "V1_17"
            and settings["patch_version"] == "V1_17"
        )
        assert not list(output.glob("*V1_15*")), "Old release name in V1_17 output"
        images = pd.read_csv(output / "image_summary.csv")
        errors = pd.read_csv(output / "processing_errors.csv")
        machine = pd.read_csv(output / "machine_comparison_detailed.csv")
        assert images.status.eq("OK").all() and errors.empty
        assert images.locator_mode_selected.eq("dark-line").all()
        assert len(machine) > 0
        psd = pd.read_csv(output / "PSD/per_structure_psd_summary.csv")
        assert set(psd.engine) == {"V10", "V13"}
        assert psd.status.eq("OK").all()
        png_count = len(list(output.rglob("*.png")))
        assert png_count > 10
        report["unpacked_full_export"] = dict(
            status="PASS",
            image_statuses=list(images.status),
            locator_modes=list(images.locator_mode_selected),
            processing_errors=len(errors),
            machine_rows=len(machine),
            psd_curves=len(psd),
            png_exports=png_count,
            mixed_cd_nm=float(images["旋转_mixed_CD_nm"].iloc[0]),
        )
        summary = pd.read_csv(output / "measurement_summary.csv")
        assert summary.columns[-1] == "旋转_pitch_CD_nm"
        assert summary.pitch_count.eq(3).all()
        assert (summary["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        report["unpacked_full_export"]["pitch_count"] = list(summary.pitch_count)
        report["unpacked_full_export"]["rotated_pitch_cd_nm"] = list(
            summary["旋转_pitch_CD_nm"]
        )
        line_output = unpacked / "line_export_smoke"
        run(
            [
                "sem_cd_measure_200k_batch_V1_17.py",
                "--root",
                "examples/input_dark_line",
                "--output",
                line_output,
                "--pattern",
                "line",
                "--pixel-size",
                "1",
                "--line-reference-nm",
                "100",
                "--no-auto-machine-comparison",
                "--skip-statistics-plots",
                "--viterbi",
                "1",
                "--erf-fit",
                "1",
            ],
            unpacked,
        )
        line = pd.read_csv(line_output / "measurement_summary.csv")
        errors = pd.read_csv(line_output / "processing_errors.csv")
        psd = pd.read_csv(line_output / "PSD/per_structure_psd_summary.csv")
        assert line.status.eq("OK").all() and errors.empty
        assert abs(line.loc[line.method == "V13", "旋转_CD_nm"].iloc[0] - 100) < 1
        assert (
            line.pitch_count.eq(3).all()
            and (line["旋转_pitch_CD_nm"] - 160).abs().lt(1).all()
        )
        assert set(psd.engine) == {"V10", "V13"} and psd.status.eq("OK").all()
        book = load_workbook(
            line_output / "CD_measurement_200K_V1_17_results.xlsx", read_only=True
        )
        assert next(book["measurement_summary"].values)[-1] == "旋转_pitch_CD_nm"
        assert (
            "line_trench_pairs" in book.sheetnames
            and "line_source_samples" in book.sheetnames
        )
        book.close()
        from check_shared_lines import verify_shared_exports

        def require(condition, reason):
            assert condition, reason

        shared = verify_shared_exports(line_output, require)
        report["unpacked_stained_line_export"] = dict(
            status="PASS",
            shared_edges=shared,
            viterbi=True,
            erf=True,
            processing_errors=len(errors),
            psd_curves=len(psd),
            rotated_pitch_cd_nm=list(line["旋转_pitch_CD_nm"]),
            pitch_count=list(line.pitch_count),
            mixed_cd_nm=float(line.loc[line.method == "mixed", "旋转_CD_nm"].iloc[0]),
        )
    evidence = ROOT / "audits/v1_17_release_evidence.json"
    evidence.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
